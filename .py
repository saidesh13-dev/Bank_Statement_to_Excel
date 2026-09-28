import argparse
import re
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

DATE_PATTERNS = (
    "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%m-%d-%Y",
    "%Y/%m/%d", "%Y-%m-%d", "%d %b %Y", "%d %B %Y",
    "%b %d %Y", "%B %d %Y",
)
DATE_RE = re.compile(
    r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}[/-]\d{1,2}[/-]\d{1,2}|"
    r"\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4}|[A-Za-z]{3,9}\s+\d{1,2},?\s+\d{4})\b",
    re.IGNORECASE,
)
SHORT_DATE_RE = re.compile(r"^\d{1,2}\s+[A-Za-z]{3,9}\b", re.IGNORECASE)
MONEY_RE = re.compile(
    r"(?<![A-Za-z0-9])(?:[+-]?\s*R\s*)?(?:\d[\d,]*\.\d{2}|"
    r"\d{1,3}(?:\s\d{3})+\.\d{2}|\d+\.\d{2})(?:\s*(?:Cr|Dr))?(?![A-Za-z0-9])",
    re.IGNORECASE,
)


def clean_line(value):
    return re.sub(r"\s+", " ", value or "").strip()


def parse_date(value, fallback_year=None):
    value = clean_line(value)
    for pattern in DATE_PATTERNS:
        try:
            return datetime.strptime(value, pattern).date()
        except ValueError:
            continue
    if fallback_year is not None:
        for pattern in ("%d %b %Y", "%d %B %Y"):
            try:
                return datetime.strptime(f"{value} {fallback_year}", pattern).date()
            except ValueError:
                continue
    return None


def parse_money(value):
    text = str(value or "").strip().replace(",", "").replace(" ", "")
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    text = text.replace("€", "").replace("$", "").replace("£", "")
    try:
        return float(text or 0)
    except ValueError:
        return 0.0


def parse_money_token(token):
    text = (token or "").strip().replace(",", "").replace(" ", "")
    if not text:
        return 0.0
    sign = -1.0 if text.startswith("-") else 1.0
    text = text.lstrip("+-")
    if text.startswith(("R", "r")):
        text = text[1:]
    suffix = text[-2:].lower()
    if suffix in {"dr", "cr"}:
        sign = -1.0 if suffix == "dr" else 1.0
        text = text[:-2]
    try:
        return sign * float(text)
    except ValueError:
        return 0.0


def extract_text_from_pdf(pdf_path):
    with pdfplumber.open(pdf_path) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def parse_pipe_row(line):
    cells = [clean_line(part) for part in line.split("|")]
    if len(cells) < 5 or not DATE_RE.search(cells[0]):
        return None
    date_value = parse_date(DATE_RE.search(cells[0]).group(0))
    if date_value is None:
        return None
    debit = parse_money(cells[2])
    credit = parse_money(cells[3])
    balance = parse_money(cells[4])
    if debit == 0 and credit == 0 and balance != 0:
        return None
    return {"Date": date_value, "Description": cells[1] or "Transaction",
            "Debit": abs(debit) if debit else 0.0,
            "Credit": abs(credit) if credit else 0.0, "Balance": balance}


def parse_text_row(line, fallback_year=None):
    date_match = DATE_RE.search(line)
    if date_match:
        date_value = parse_date(date_match.group(0), fallback_year)
        after_date = line[date_match.end():].strip()
    else:
        short_match = re.search(r"\b\d{1,2}\s+[A-Za-z]{3,9}\b", line)
        if fallback_year is None or short_match is None:
            return None
        date_value = parse_date(short_match.group(0), fallback_year)
        after_date = line[short_match.end():].strip()
    if date_value is None or not after_date:
        return None

    matches = []
    for match in MONEY_RE.finditer(after_date):
        token = match.group(0)
        value = parse_money_token(token)
        if re.search(r"(?:R|Cr|Dr|\d[.,]\d)", token, re.IGNORECASE):
            matches.append((match, value))
    if len(matches) < 2:
        return None

    short_date_row = SHORT_DATE_RE.match(line) is not None
    amount_match, amount = matches[0] if short_date_row else matches[-2]
    balance_match, balance = matches[1] if short_date_row else matches[-1]
    amount_token = amount_match.group(0)
    description = after_date[:amount_match.start()].strip(" -:;|") or "Transaction"
    lower_description = description.lower()
    if abs(amount) <= 0.01 and any(word in lower_description for word in ["proof", "pmt", "email", "sms"]):
        return None

    is_credit = bool(re.search(r"cr\s*$", amount_token, re.IGNORECASE))
    is_credit = is_credit or bool(re.search(r"\b(credit|deposit|salary|refund|received|incoming)\b", lower_description))
    if amount < 0:
        debit, credit = abs(amount), 0.0
    elif is_credit:
        debit, credit = 0.0, abs(amount)
    elif short_date_row:
        debit, credit = abs(amount), 0.0
    else:
        debit, credit = 0.0, abs(amount)
    return {"Date": date_value, "Description": description, "Debit": debit,
            "Credit": credit, "Balance": balance}


def parse_bank_statement(pdf_path):
    text = extract_text_from_pdf(pdf_path)
    statement_year = None
    for line in text.splitlines():
        match = re.search(r"\b(19\d{2}|20\d{2})\b", line)
        if match and "statement" in line.lower():
            statement_year = int(match.group(1))

    records = []
    for raw_line in text.splitlines():
        line = clean_line(raw_line)
        if not line:
            continue
        record = parse_pipe_row(line) if "|" in line else None
        record = record or parse_text_row(line, statement_year)
        if record and record["Description"].lower() not in {"date", "description", "debit", "credit", "balance", "transaction"}:
            records.append(record)
    if not records:
        raise ValueError("No valid bank transactions were found in the PDF.")

    records.sort(key=lambda item: item["Date"])
    running_balance = 0.0
    normalized = []
    for record in records:
        balance = record.get("Balance", 0.0)
        if balance:
            running_balance = float(balance)
        else:
            running_balance += record["Credit"] - record["Debit"]
            balance = running_balance
        normalized.append({"Date": record["Date"], "Description": record["Description"],
                           "Debit": round(record["Debit"], 2), "Credit": round(record["Credit"], 2),
                           "Balance": round(balance, 2)})
    return normalized


def export_to_excel(records, output_excel_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bank Statement"
    sheet.append(["Date", "Description", "Debit", "Credit", "Balance"])
    for item in records:
        sheet.append([item["Date"].strftime("%d/%m/%Y"), item["Description"], item["Debit"], item["Credit"], item["Balance"]])
    border = Border(*(Side(style="thin", color="D9D9D9") for _ in range(4)))
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(bold=True)
        cell.border = border
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.border = border
    for row in sheet.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = '[$$-en-US]#,##0.00'
    for column_cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells) + 2
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
    sheet.freeze_panes = "A2"
    last_row = sheet.max_row
    sheet["F2"], sheet["G2"] = "Total Debits", f"=SUM(C2:C{last_row})"
    sheet["F3"], sheet["G3"] = "Total Credits", f"=SUM(D2:D{last_row})"
    sheet["F4"], sheet["G4"] = "Net Movement", f"=SUM(D2:D{last_row})-SUM(C2:C{last_row})"
    for row in (2, 3, 4):
        sheet[f"F{row}"].font = Font(bold=True)
        sheet[f"G{row}"].number_format = '[$$-en-US]#,##0.00'
    output_path = Path(output_excel_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def read_bank_statement(pdf_path, output_excel_path):
    records = parse_bank_statement(pdf_path)
    return records, export_to_excel(records, output_excel_path)


def upload_and_export(uploaded_pdf, output_name="bank_statement.xlsx"):
    with NamedTemporaryFile(suffix=".pdf", delete=False) as temp_pdf:
        temp_pdf.write(uploaded_pdf.read())
        temp_path = temp_pdf.name
    try:
        records, output_path = read_bank_statement(temp_path, output_name)
        return records, str(output_path)
    finally:
        Path(temp_path).unlink(missing_ok=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("pdf_path")
    parser.add_argument("output_excel_path", nargs="?", default="bank_statement.xlsx")
    args = parser.parse_args()
    read_bank_statement(args.pdf_path, args.output_excel_path)
